"""End to end over real files: ingest a Word doc and a PDF, then answer with a faked LLM."""

import json

import docx
import httpx
import openai
import pymupdf
import pytest

import ingest
import rag


@pytest.fixture
def library(tmp_path, monkeypatch):
    docs = tmp_path / "docs"
    (docs / "Policy").mkdir(parents=True)

    d = docx.Document()
    d.add_heading("Overnight stays", level=1)
    d.add_paragraph("12. Cadets may sleep overnight in Army Reserve Centres, provided that the OC "
                    "has written approval from the Wing.")
    d.add_paragraph("13. Adult staff of both sexes must be present when mixed groups stay overnight.")
    d.add_heading("Transport", level=1)
    d.add_paragraph("14. Minibus drivers must hold a D1 licence.")
    d.save(docs / "Policy" / "ACP 20 Overnight.docx")

    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 100), "Welcome to the uniform guide.")
    page = pdf.new_page()
    page.insert_text((50, 100), "4.2")
    page.insert_text((100, 100), "Berets are worn at all times outdoors except when flying.")
    pdf.save(docs / "Uniform Guide.pdf")

    monkeypatch.setattr(rag, "DOCS_DIR", docs)
    monkeypatch.setattr(ingest, "STATE_FILE", tmp_path / "state.json")
    ingest.main()
    return docs


def fake_llm(monkeypatch, reply):
    """Answers every question with reply(messages), recording what was sent."""
    sent = []

    def chat(messages, retries=5):
        sent.append(messages)
        return reply(messages) if callable(reply) else reply

    monkeypatch.setattr(rag, "_chat", chat)
    return sent


def source_number(messages, filename):
    """The [n] the prompt gave a file, so the fake reply can cite it whatever the retrieval order."""
    for block in messages[1]["content"].split("\n\n---\n\n"):
        head = block.removeprefix("Sources:\n\n")
        if filename in head.split("\n")[0]:
            return int(head[1:head.index("]")])
    raise AssertionError(f"{filename} not in prompt")


def test_an_answer_cites_the_exact_lines_with_page_and_paragraph(library, monkeypatch):
    def reply(messages):
        w, p = source_number(messages, "ACP 20"), source_number(messages, "Uniform")
        return json.dumps({
            "answer": f"Berets are worn outdoors [{p}]. Cadets may stay overnight with Wing approval [{w}].",
            "citations": [
                {"source": p, "quote": "Berets are worn at all times outdoors except when flying."},
                {"source": w, "quote": "provided that the OC has written approval from the Wing"},
            ],
        })

    fake_llm(monkeypatch, reply)
    res = rag.answer("Can cadets sleep overnight and when are berets worn?")

    assert res["found"] is True
    assert res["answer"] == "Berets are worn outdoors [1]. Cadets may stay overnight with Wing approval [2]."
    assert [s["n"] for s in res["sources"]] == [1, 2]
    uniform, acp = res["sources"]

    assert uniform["filename"] == "Uniform Guide.pdf"
    assert uniform["location"] == "page 2, para 4.2"
    assert uniform["url"].endswith("/Uniform%20Guide.pdf#page=2")
    [q] = uniform["quotes"]
    assert q["text"] == "Berets are worn at all times outdoors except when flying."
    assert uniform["text"][q["start"]:q["end"]] == q["text"]
    assert (q["page"], q["location"]) == (2, "page 2, para 4.2")

    assert acp["location"] == "Overnight stays, para 12"
    assert acp["url"].endswith("/Policy/ACP%2020%20Overnight.docx")  # no #page for Word
    assert acp["quotes"][0]["text"] == "provided that the OC has written approval from the Wing"
    assert acp["quotes"][0]["page"] is None
    assert res["related"] == []


def test_the_prompt_carries_numbered_sources_and_the_question(library, monkeypatch):
    sent = fake_llm(monkeypatch, '{"answer": "NOT_FOUND", "citations": []}')
    rag.answer("Who signs off minibus drivers?")
    system, user = sent[0][0]["content"], sent[0][1]["content"]
    assert "JSON" in system  # providers' JSON mode requires the word in the prompt
    assert user.startswith("Sources:\n\n[1] ") and user.endswith("Question: Who signs off minibus drivers?")
    assert "D1 licence" in user


def test_sources_the_answer_never_cites_are_left_out(library, monkeypatch):
    def reply(messages):
        w = source_number(messages, "ACP 20")
        return json.dumps({"answer": f"Drivers need a D1 licence [{w}].",
                           "citations": [{"source": w, "quote": "Minibus drivers must hold a D1 licence."}]})

    fake_llm(monkeypatch, reply)
    res = rag.answer("What licence do minibus drivers need?")
    assert [s["filename"] for s in res["sources"]] == ["ACP 20 Overnight.docx"]
    assert res["sources"][0]["location"] == "Transport, para 14"


def test_a_made_up_quote_is_dropped_but_the_cited_source_stays(library, monkeypatch):
    def reply(messages):
        w = source_number(messages, "ACP 20")
        made_up = "Only commissioned officers may ever drive a minibus."
        return json.dumps({"answer": f"Only officers may drive [{w}].",
                           "citations": [{"source": w, "quote": made_up}]})

    fake_llm(monkeypatch, reply)
    [src] = rag.answer("Who may drive?")["sources"]
    assert src["quotes"] == []
    assert src["location"]  # falls back to where the chunk starts


def test_citations_of_sources_that_were_never_given_are_ignored(library, monkeypatch):
    fake_llm(monkeypatch, json.dumps({"answer": "Yes [42].",
                                      "citations": [{"source": 42, "quote": "anything at all here"}]}))
    res = rag.answer("Can cadets sleep overnight?")
    assert res["answer"] == "Yes."
    assert res["sources"] == []
    assert res["related"]  # nothing to cite, so point at the closest documents


def test_a_plain_text_reply_still_works_without_quotes(library, monkeypatch):
    fake_llm(monkeypatch, "Cadets may stay overnight [1].")
    res = rag.answer("Can cadets sleep overnight?")
    assert res["answer"] == "Cadets may stay overnight [1]."
    assert res["sources"][0]["quotes"] == []


@pytest.mark.parametrize("reply", [
    '{"answer": "NOT_FOUND", "citations": []}',
    "I couldn't find that in the documents.",
    "I couldn’t find that in the documents.",
    '{"answer": "", "citations": []}',
])
def test_no_answer_offers_the_closest_documents_instead_of_sources(library, monkeypatch, reply):
    fake_llm(monkeypatch, reply)
    res = rag.answer("What is the canteen's opening time?")
    assert res["answer"] == rag.NOT_FOUND
    assert res["found"] is False and res["sources"] == []
    assert {r["filename"] for r in res["related"]} == {"ACP 20 Overnight.docx", "Uniform Guide.pdf"}


def test_an_empty_index_answers_not_found_without_calling_the_llm(monkeypatch):
    sent = fake_llm(monkeypatch, "should not be asked")
    assert rag.answer("anything") == {"answer": rag.NOT_FOUND, "found": False, "sources": [], "related": []}
    assert sent == []


def _response(status):
    return httpx.Response(status, request=httpx.Request("POST", "https://llm.test/v1/chat/completions"))


def test_a_rate_limited_provider_gives_a_try_again_message_and_related_docs(library, monkeypatch):
    def chat(messages, retries=5):
        raise openai.RateLimitError("slow down", response=_response(429), body=None)

    monkeypatch.setattr(rag, "_chat", chat)
    res = rag.answer("Can cadets sleep overnight?")
    assert "Try again in a minute" in res["answer"]
    assert res["found"] is False and res["related"]


def test_a_missing_api_key_is_reported_not_raised(library):
    res = rag.answer("Can cadets sleep overnight?")
    assert res["found"] is False
    assert "couldn't answer" in res["answer"]


class FakeCompletions:
    def __init__(self, outcomes):
        self.outcomes, self.calls = list(outcomes), []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        out = self.outcomes.pop(0)
        if isinstance(out, Exception):
            raise out
        msg = type("M", (), {"content": out})
        return type("R", (), {"choices": [type("C", (), {"message": msg})]})


def fake_client(monkeypatch, outcomes):
    completions = FakeCompletions(outcomes)
    client = type("Client", (), {"chat": type("Chat", (), {"completions": completions})})
    monkeypatch.setattr(rag, "OpenAI", lambda **_: client)
    monkeypatch.setattr(rag.time, "sleep", lambda s: None)
    monkeypatch.setenv("LLM_API_KEY", "test")
    return completions


def test_chat_asks_for_json_mode(monkeypatch):
    calls = fake_client(monkeypatch, ['{"answer": "x"}'])
    assert rag._chat([]) == '{"answer": "x"}'
    assert calls.calls[0]["response_format"] == {"type": "json_object"}


def test_a_provider_that_rejects_json_mode_is_retried_without_it(monkeypatch):
    bad = openai.BadRequestError("json mode unsupported", response=_response(400), body=None)
    calls = fake_client(monkeypatch, [bad, "plain"])
    assert rag._chat([]) == "plain"
    assert "response_format" not in calls.calls[1]


def test_a_second_bad_request_is_not_swallowed(monkeypatch):
    bad = openai.BadRequestError("bad", response=_response(400), body=None)
    fake_client(monkeypatch, [bad, bad])
    with pytest.raises(openai.BadRequestError):
        rag._chat([])


def test_rate_limits_are_retried_then_raised(monkeypatch):
    limited = openai.RateLimitError("429", response=_response(429), body=None)
    calls = fake_client(monkeypatch, [limited, "ok"])
    assert rag._chat([]) == "ok"
    fake_client(monkeypatch, [limited] * 3)
    with pytest.raises(openai.RateLimitError):
        rag._chat([], retries=3)
    assert len(calls.calls) == 2


def test_the_reranker_reorders_candidates_and_keeps_the_best(library, monkeypatch):
    class Reranker:
        def predict(self, pairs, **_):
            # Prefer the PDF whatever the embedding said; it sees the title-prefixed passage.
            return [10.0 if "Uniform Guide" in p else 0.0 for _, p in pairs]

    monkeypatch.setattr(rag, "get_reranker", lambda: Reranker())
    chunks = rag.retrieve("overnight stays in reserve centres", k=1)
    assert [c["meta"]["filename"] for c in chunks] == ["Uniform Guide.pdf"]
    assert chunks[0]["score"] == 10.0


def test_an_index_from_before_spans_still_answers_with_chunk_locations(monkeypatch):
    rag.get_collection().add(ids=["old::0"], documents=["Berets are worn outdoors at all times."],
                             embeddings=rag.embed(["Berets are worn outdoors at all times."]),
                             metadatas=[{"path": "Old.pdf", "filename": "Old.pdf", "chunk": 0, "page": 3}])
    fake_llm(monkeypatch, json.dumps({"answer": "Outdoors [1].",
                                      "citations": [{"source": 1, "quote": "Berets are worn outdoors at all times."}]}))
    [src] = rag.answer("When are berets worn?")["sources"]
    assert src["location"] == "page 3"
    assert src["url"].endswith("Old.pdf#page=3")
    assert src["quotes"][0]["page"] == 3


def test_reingest_skips_unchanged_files_but_redoes_an_old_index_version(library, monkeypatch):
    state_file = ingest.STATE_FILE
    state = json.loads(state_file.read_text())
    assert all(s["v"] == ingest.INDEX_VERSION for s in state.values())

    embedded = []
    real_embed = rag.embed
    monkeypatch.setattr(rag, "embed",
                        lambda texts, query=False: embedded.append(len(texts)) or real_embed(texts, query))
    ingest.main()
    assert embedded == []  # nothing changed

    for s in state.values():
        s["v"] = 1
    state_file.write_text(json.dumps(state))
    ingest.main()
    assert len(embedded) == 2  # both files redone


def test_the_ask_endpoint_validates_and_answers(library, monkeypatch):
    from fastapi.testclient import TestClient

    import app

    fake_llm(monkeypatch, '{"answer": "NOT_FOUND", "citations": []}')
    client = TestClient(app.app)
    assert client.post("/ask", json={"question": ""}).status_code == 422
    res = client.post("/ask", json={"question": "  anything?  "})
    assert res.status_code == 200 and res.json()["found"] is False
    assert client.get("/health").json()["chunks"] > 0


def test_the_cli_prints_quotes_and_closest_documents(capsys):
    from cli import print_result

    print_result({"answer": "Yes [1].", "sources": [{
        "n": 1, "filename": "A.pdf", "location": "page 2, para 4", "url": "u#page=2", "score": 0.5,
        "quotes": [{"location": "page 2, para 4", "text": "Berets are\nworn."}]}], "related": []})
    out = capsys.readouterr().out
    assert '> page 2, para 4: "Berets are worn."' in out

    print_result({"answer": rag.NOT_FOUND, "sources": [],
                  "related": [{"filename": "B.docx", "location": "", "url": "v"}]})
    assert "Closest documents:\n  - B.docx  v" in capsys.readouterr().out
